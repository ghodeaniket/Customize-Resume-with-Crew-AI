"""
Resume export utilities.

This module provides functions for exporting resumes in different formats,
including plain text, PDF, and DOCX.
"""

import streamlit as st
import base64
from typing import Dict, Any, Optional, Tuple, List
import tempfile
import os
import io

# For PDF generation
try:
    from reportlab.lib.pagesizes import LETTER, A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False

# For DOCX generation
try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


def create_download_link(file_content: bytes, filename: str, mime_type: str, button_text: str) -> None:
    """
    Create a download button for a file.
    
    Args:
        file_content: The file content as bytes
        filename: The filename
        mime_type: The MIME type of the file
        button_text: Text to display on the button
    """
    b64 = base64.b64encode(file_content).decode()
    href = f'data:{mime_type};base64,{b64}'
    
    st.download_button(
        label=button_text,
        data=file_content,
        file_name=filename,
        mime=mime_type,
        use_container_width=True
    )


def export_as_text(text_content: str, filename: str = "resume.txt") -> None:
    """
    Create a download button for plain text content.
    
    Args:
        text_content: The text content
        filename: The filename to use
    """
    create_download_link(
        text_content.encode('utf-8'),
        filename,
        "text/plain",
        f"Download as Text (.txt)"
    )


def export_as_html(text_content: str, title: str = "Resume", filename: str = "resume.html") -> None:
    """
    Create a download button for HTML content.
    
    Args:
        text_content: The text content
        title: The HTML document title
        filename: The filename to use
    """
    # Convert plain text to HTML with proper line breaks
    html_content = text_content.replace('\n', '<br>\n')
    
    # Create a simple HTML document
    html_doc = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{title}</title>
    <style>
        body {{
            font-family: Arial, sans-serif;
            line-height: 1.6;
            margin: 2rem auto;
            max-width: 800px;
            padding: 0 1rem;
        }}
        .resume-container {{
            border: 1px solid #e2e8f0;
            padding: 2rem;
            box-shadow: 0 4px 6px rgba(0, 0, 0, 0.1);
        }}
        @media print {{
            body {{
                margin: 0;
                padding: 0;
            }}
            .resume-container {{
                border: none;
                box-shadow: none;
                padding: 1rem;
            }}
        }}
    </style>
</head>
<body>
    <div class="resume-container">
        {html_content}
    </div>
</body>
</html>
"""
    
    # Create download button
    create_download_link(
        html_doc.encode('utf-8'),
        filename,
        "text/html",
        f"Download as HTML (.html)"
    )


def export_as_pdf(text_content: str, filename: str = "resume.pdf") -> Optional[bytes]:
    """
    Export resume as PDF using ReportLab.
    
    Args:
        text_content: The text content
        filename: The filename to use
        
    Returns:
        Optional[bytes]: The PDF content as bytes, or None if ReportLab is not available
    """
    # Check if ReportLab is available
    if not REPORTLAB_AVAILABLE:
        st.warning("PDF export is not available. ReportLab library is required.")
        return None
    
    # Create a buffer to store the PDF data
    buffer = io.BytesIO()
    
    # Create the PDF document
    doc = SimpleDocTemplate(
        buffer,
        pagesize=LETTER,
        rightMargin=72,
        leftMargin=72,
        topMargin=72,
        bottomMargin=72
    )
    
    # Create styles
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name='ResumeNormal',
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        spaceAfter=10
    ))
    styles.add(ParagraphStyle(
        name='ResumeHeading',
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        spaceAfter=12
    ))
    
    # Split text into lines and process them
    lines = text_content.splitlines()
    flowables = []
    
    # Process lines and build PDF content
    current_paragraph = ""
    
    for line in lines:
        line = line.strip()
        
        # Check if this is a potential heading (uppercase, short)
        if line.isupper() and len(line) <= 30 and line:
            # If we have accumulated content, add it as a paragraph
            if current_paragraph:
                flowables.append(Paragraph(current_paragraph, styles['ResumeNormal']))
                current_paragraph = ""
            
            # Add this line as a heading
            flowables.append(Paragraph(line, styles['ResumeHeading']))
        else:
            # If line is empty and we have content, end the paragraph
            if not line and current_paragraph:
                flowables.append(Paragraph(current_paragraph, styles['ResumeNormal']))
                current_paragraph = ""
                flowables.append(Spacer(1, 6))
            # If line is not empty, add it to the current paragraph
            elif line:
                if current_paragraph:
                    current_paragraph += "<br/>" + line
                else:
                    current_paragraph = line
    
    # Don't forget the last paragraph
    if current_paragraph:
        flowables.append(Paragraph(current_paragraph, styles['ResumeNormal']))
    
    # Build the PDF
    doc.build(flowables)
    
    # Get the PDF data
    pdf_data = buffer.getvalue()
    buffer.close()
    
    # Create the download button
    create_download_link(
        pdf_data,
        filename,
        "application/pdf",
        f"Download as PDF (.pdf)"
    )
    
    return pdf_data


def export_as_docx(text_content: str, filename: str = "resume.docx") -> Optional[bytes]:
    """
    Export resume as DOCX using python-docx.
    
    Args:
        text_content: The text content
        filename: The filename to use
        
    Returns:
        Optional[bytes]: The DOCX content as bytes, or None if python-docx is not available
    """
    # Check if python-docx is available
    if not DOCX_AVAILABLE:
        st.warning("DOCX export is not available. python-docx library is required.")
        return None
    
    # Create a Document
    doc = Document()
    
    # Set document properties
    doc.core_properties.title = "Resume"
    
    # Set default font
    style = doc.styles['Normal']
    style.font.size = Pt(11)
    style.font.name = 'Calibri'
    
    # Set margins (1 inch)
    section = doc.sections[0]
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    
    # Split text into lines and process them
    lines = text_content.splitlines()
    current_paragraph = None
    
    for line in lines:
        line_stripped = line.strip()
        
        # Check if this is a potential heading (uppercase, short)
        if line_stripped.isupper() and len(line_stripped) <= 30 and line_stripped:
            # Add this line as a heading
            heading = doc.add_paragraph(line_stripped)
            heading.style = 'Heading 2'
            heading.space_after = Pt(8)
        elif not line_stripped:
            # Empty line - add a paragraph break
            if current_paragraph:
                current_paragraph = None
        else:
            # Regular content - add to paragraph
            if not current_paragraph:
                current_paragraph = doc.add_paragraph()
            
            # Add text with a line break if not the first line in the paragraph
            if current_paragraph.text:
                current_paragraph.add_run('\n' + line)
            else:
                current_paragraph.add_run(line)
    
    # Save the document to a bytes buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    
    docx_data = buffer.getvalue()
    buffer.close()
    
    # Create the download button
    create_download_link(
        docx_data,
        filename,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        f"Download as Word (.docx)"
    )
    
    return docx_data


def copy_to_clipboard(text_content: str) -> None:
    """
    Create a button to copy text content to clipboard.
    
    Args:
        text_content: The text content to copy
    """
    # Streamlit doesn't have direct clipboard access, so we use a JavaScript snippet
    # with st.components.v1.html or st.markdown with unsafe_allow_html=True
    
    # Encode the text content for the JavaScript function
    encoded_text = text_content.replace('\n', '\\n').replace('\"', '\\"')
    
    # Create a unique ID for the button
    button_id = f"copy-button-{hash(text_content) % 10000}"
    
    # JavaScript function to copy text to clipboard
    js_code = f"""
    <button id="{button_id}" style="
        background-color: #f1f5f9;
        border: 1px solid #cbd5e1;
        padding: 0.5rem 1rem;
        border-radius: 0.25rem;
        font-size: 0.875rem;
        color: #334155;
        cursor: pointer;
        transition: all 0.2s;
        width: 100%;
        text-align: center;
        margin: 0.5rem 0;
    ">
        Copy to Clipboard
    </button>
    
    <script>
    const copyButton{button_id.replace('-', '_')} = document.getElementById('{button_id}');
    copyButton{button_id.replace('-', '_')}.addEventListener('click', function() {{
        // Create a temporary textarea element
        const textarea = document.createElement('textarea');
        textarea.value = "{encoded_text}";
        textarea.setAttribute('readonly', '');
        textarea.style.position = 'absolute';
        textarea.style.left = '-9999px';
        document.body.appendChild(textarea);
        
        // Select and copy the text
        textarea.select();
        document.execCommand('copy');
        
        // Clean up
        document.body.removeChild(textarea);
        
        // Update button text temporarily
        const originalText = this.innerText;
        this.innerText = 'Copied!';
        this.style.backgroundColor = '#d1fae5';
        this.style.color = '#065f46';
        this.style.borderColor = '#6ee7b7';
        
        // Revert button text after 2 seconds
        setTimeout(() => {{
            this.innerText = originalText;
            this.style.backgroundColor = '#f1f5f9';
            this.style.color = '#334155';
            this.style.borderColor = '#cbd5e1';
        }}, 2000);
    }});
    </script>
    """
    
    st.markdown(js_code, unsafe_allow_html=True)


def get_print_friendly_view(text_content: str) -> str:
    """
    Generate a print-friendly HTML view of the resume.
    
    Args:
        text_content: The resume text content
        
    Returns:
        str: The HTML code for a print-friendly view
    """
    # Convert plain text to HTML with proper line breaks
    html_content = text_content.replace('\n', '<br>\n')
    
    # Create a print-friendly HTML document with print styles
    html_doc = f"""
    <div class="print-container">
        <style>
            @media print {{
                body {{ 
                    font-family: 'Helvetica', 'Arial', sans-serif;
                    line-height: 1.6;
                    margin: 0;
                    padding: 0;
                    color: #000;
                }}
                .print-container {{
                    padding: 0.5in;
                }}
                .print-button {{
                    display: none;
                }}
            }}
            .print-container {{
                font-family: 'Helvetica', 'Arial', sans-serif;
                line-height: 1.6;
                padding: 1rem;
                background-color: white;
                border: 1px solid #e2e8f0;
                box-shadow: 0 4px 6px rgba(0, 0, 0, 0.1);
            }}
            .print-button {{
                background-color: #3b82f6;
                color: white;
                border: none;
                padding: 0.5rem 1rem;
                font-size: 1rem;
                border-radius: 0.25rem;
                cursor: pointer;
                margin: 1rem 0;
                display: block;
            }}
        </style>
        <button class="print-button" onclick="window.print()">Print Resume</button>
        <div class="resume-content">
            {html_content}
        </div>
    </div>
    """
    
    return html_doc
