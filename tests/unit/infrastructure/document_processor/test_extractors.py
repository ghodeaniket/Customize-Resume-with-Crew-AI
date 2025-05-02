"""Tests for document extractors."""
import pytest
from unittest.mock import patch, MagicMock, AsyncMock
import io
import docx

from app.core.exceptions import DocumentProcessingError
from app.infrastructure.document_processor.pdf_extractor import PDFExtractor
from app.infrastructure.document_processor.docx_extractor import DOCXExtractor
from app.infrastructure.document_processor.txt_extractor import TXTExtractor


class TestPDFExtractor:
    """Test cases for PDFExtractor."""
    
    @pytest.fixture
    def extractor(self):
        """Create PDF extractor for testing."""
        return PDFExtractor(cache_enabled=False)
    
    @pytest.fixture
    def pdf_content(self):
        """Sample PDF content."""
        return b"%PDF-1.5\nsome pdf content"
    
    @pytest.mark.asyncio
    @patch("app.infrastructure.document_processor.pdf_extractor.PDFExtractor._extract_with_pypdf2")
    async def test_extract_text(self, mock_extract, extractor, pdf_content):
        """Test PDF extraction."""
        mock_extract.return_value = "Extracted PDF content"
        
        result = await extractor.extract_text(pdf_content, "test.pdf")
        
        assert result == "Extracted PDF content"
        mock_extract.assert_called_once()
    
    @pytest.mark.asyncio
    @patch("app.infrastructure.document_processor.pdf_extractor.PYMUPDF_AVAILABLE", True)
    @patch("app.infrastructure.document_processor.pdf_extractor.PDFExtractor._extract_with_pymupdf")
    @patch("app.infrastructure.document_processor.pdf_extractor.PDFExtractor._extract_with_pypdf2")
    async def test_extract_with_fallback(self, mock_pypdf2, mock_pymupdf, extractor, pdf_content):
        """Test PDF extraction with fallback."""
        # PyMuPDF fails, PyPDF2 succeeds
        mock_pymupdf.side_effect = Exception("PyMuPDF error")
        mock_pypdf2.return_value = "Extracted with PyPDF2"
        
        result = await extractor.extract_text(pdf_content, "test.pdf")
        
        assert result == "Extracted with PyPDF2"
        mock_pymupdf.assert_called_once()
        mock_pypdf2.assert_called_once()
    
    @pytest.mark.asyncio
    @patch("app.infrastructure.document_processor.pdf_extractor.PYMUPDF_AVAILABLE", True)
    @patch("app.infrastructure.document_processor.pdf_extractor.PDFExtractor._extract_with_pymupdf")
    @patch("app.infrastructure.document_processor.pdf_extractor.PDFExtractor._extract_with_pypdf2")
    async def test_extract_all_methods_fail(self, mock_pypdf2, mock_pymupdf, extractor, pdf_content):
        """Test PDF extraction when all methods fail."""
        mock_pymupdf.side_effect = Exception("PyMuPDF error")
        mock_pypdf2.side_effect = Exception("PyPDF2 error")
        
        with pytest.raises(DocumentProcessingError, match="All PDF extraction methods failed"):
            await extractor.extract_text(pdf_content, "test.pdf")


class TestDOCXExtractor:
    """Test cases for DOCXExtractor."""
    
    @pytest.fixture
    def extractor(self):
        """Create DOCX extractor for testing."""
        return DOCXExtractor(cache_enabled=False)
    
    @pytest.fixture
    def docx_content(self):
        """Create a simple DOCX file."""
        doc = docx.Document()
        doc.add_paragraph("Test paragraph 1")
        doc.add_paragraph("Test paragraph 2")
        
        # Create a table
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Cell 1"
        table.cell(0, 1).text = "Cell 2"
        table.cell(1, 0).text = "Cell 3"
        table.cell(1, 1).text = "Cell 4"
        
        # Save to bytes
        docx_bytes = io.BytesIO()
        doc.save(docx_bytes)
        docx_bytes.seek(0)
        return docx_bytes.read()
    
    @pytest.mark.asyncio
    @patch("asyncio.to_thread")
    async def test_extract_text(self, mock_to_thread, extractor, docx_content):
        """Test DOCX extraction."""
        mock_to_thread.return_value = "Test paragraph 1\n\nTest paragraph 2\n\nCell 1 | Cell 2\nCell 3 | Cell 4"
        
        result = await extractor.extract_text(docx_content, "test.docx")
        
        assert result == "Test paragraph 1\n\nTest paragraph 2\n\nCell 1 | Cell 2\nCell 3 | Cell 4"
        mock_to_thread.assert_called_once()
    
    @pytest.mark.asyncio
    @patch("asyncio.to_thread")
    async def test_extract_docx_error(self, mock_to_thread, extractor, docx_content):
        """Test DOCX extraction with error."""
        mock_to_thread.side_effect = Exception("DOCX error")
        
        with pytest.raises(DocumentProcessingError, match="DOCX extraction failed"):
            await extractor.extract_text(docx_content, "test.docx")


class TestTXTExtractor:
    """Test cases for TXTExtractor."""
    
    @pytest.fixture
    def extractor(self):
        """Create TXT extractor for testing."""
        return TXTExtractor(cache_enabled=False)
    
    @pytest.mark.asyncio
    async def test_extract_text(self, extractor):
        """Test text extraction."""
        content = "This is a plain text document.".encode('utf-8')
        
        result = await extractor.extract_text(content, "test.txt")
        
        assert result == "This is a plain text document."
    
    @pytest.mark.asyncio
    async def test_extract_utf8_with_replacement(self, extractor):
        """Test text extraction with UTF-8 replacement."""
        # Create a byte string with invalid UTF-8 sequence
        content = b"This has an invalid \x80 UTF-8 byte."
        
        result = await extractor.extract_text(content, "test.txt")
        
        # Should contain replacement character
        assert "This has an invalid" in result
        assert "UTF-8 byte" in result
    
    @pytest.mark.asyncio
    async def test_extract_text_error(self, extractor):
        """Test text extraction with error."""
        # Disable fallback to trigger error
        extractor.fallback_encoding = False
        
        # Create a byte string with invalid UTF-8 sequence
        content = b"This has an invalid \x80 UTF-8 byte."
        
        with pytest.raises(DocumentProcessingError, match="Failed to decode text"):
            await extractor.extract_text(content, "test.txt")
