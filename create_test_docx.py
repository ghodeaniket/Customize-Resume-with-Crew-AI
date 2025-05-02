#!/usr/bin/env python3
"""
Create Test DOCX Resume

This script generates a sample DOCX resume for testing the Resume Customizer application.
"""
import os
import sys
from pathlib import Path

try:
    # Try to import docx
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
except ImportError:
    print("ERROR: The 'python-docx' package is required to create DOCX files.")
    print("Please install it using: pip install python-docx")
    sys.exit(1)

# Create test_data directory if it doesn't exist
test_data_dir = Path("test_data")
test_data_dir.mkdir(exist_ok=True)

# Path for the output DOCX file
docx_path = test_data_dir / "test_resume.docx"

print(f"Creating test DOCX resume at: {docx_path}")

# Create a new Document
doc = Document()

# Set document margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Title and contact information
title = doc.add_paragraph()
title_run = title.add_run("John Doe")
title_run.bold = True
title_run.font.size = Pt(18)

role = doc.add_paragraph()
role_run = role.add_run("Software Engineer")
role_run.bold = True
role_run.font.size = Pt(14)

contact = doc.add_paragraph()
contact.add_run("john.doe@example.com | (555) 123-4567 | linkedin.com/in/johndoe")

# Summary
doc.add_heading("Summary", level=2)
summary = doc.add_paragraph("Experienced software engineer with 5+ years in backend development specializing in Python, FastAPI, and cloud technologies.")

# Skills
doc.add_heading("Skills", level=2)
skills = doc.add_paragraph()
skills.add_run("• Languages: ").bold = True
skills.add_run("Python, JavaScript, SQL, TypeScript\n")
skills.add_run("• Frameworks: ").bold = True
skills.add_run("FastAPI, Django, Flask, React\n")
skills.add_run("• Cloud: ").bold = True
skills.add_run("AWS, Docker, Kubernetes\n")
skills.add_run("• Databases: ").bold = True
skills.add_run("PostgreSQL, MongoDB\n")
skills.add_run("• AI/ML: ").bold = True
skills.add_run("CrewAI, LangChain, TensorFlow")

# Experience
doc.add_heading("Experience", level=2)

job1 = doc.add_paragraph()
job1_title = job1.add_run("Senior Software Engineer | TechCorp Inc. | 2022-Present")
job1_title.bold = True

exp1 = doc.add_paragraph()
exp1.style = 'List Bullet'
exp1.add_run("Developed and maintained microservices using FastAPI and Python")

exp2 = doc.add_paragraph()
exp2.style = 'List Bullet'
exp2.add_run("Implemented CI/CD pipelines reducing deployment time by 40%")

exp3 = doc.add_paragraph()
exp3.style = 'List Bullet'
exp3.add_run("Led a team of 3 developers in building a document processing system")

job2 = doc.add_paragraph()
job2_title = job2.add_run("Software Engineer | Innovate Solutions | 2020-2022")
job2_title.bold = True

exp4 = doc.add_paragraph()
exp4.style = 'List Bullet'
exp4.add_run("Built RESTful APIs using Flask and PostgreSQL")

exp5 = doc.add_paragraph()
exp5.style = 'List Bullet'
exp5.add_run("Migrated monolithic application to microservices architecture")

exp6 = doc.add_paragraph()
exp6.style = 'List Bullet'
exp6.add_run("Optimized database queries resulting in 30% performance improvement")

# Education
doc.add_heading("Education", level=2)
edu = doc.add_paragraph()
edu_title = edu.add_run("Bachelor of Science in Computer Science")
edu_title.bold = True
doc.add_paragraph("University of Technology, 2020")

# Save the document
doc.save(docx_path)

print(f"Successfully created test DOCX resume at: {docx_path}")
print("You can now use this file to test the Resume Customizer application.")
print("\nExample command:")
print(f"curl -X POST -F \"resume=@{docx_path}\" http://127.0.0.1:8000/api/resumes/upload")
